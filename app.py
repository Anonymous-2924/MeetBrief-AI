
import os
import io
import json
import tempfile
from datetime import datetime

import pandas as pd
import streamlit as st
from dotenv import load_dotenv
from groq import Groq
from docx import Document
from reportlab.lib import colors
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.lib.units import inch
from reportlab.platypus import (
    SimpleDocTemplate,
    Paragraph,
    Spacer,
    Table,
    TableStyle,
)

load_dotenv()

# =========================================================
# CONFIGURATION
# =========================================================

APP_TITLE = "MeetBrief AI"
TEXT_MODEL = os.getenv("GROQ_MODEL", "openai/gpt-oss-20b")
VISION_MODEL = "qwen/qwen3.8-27b"
AUDIO_MODEL = "whisper-large-v3"

st.set_page_config(
    page_title=APP_TITLE,
    page_icon="📝",
    layout="wide",
    initial_sidebar_state="expanded",
)

st.markdown(
    """
    <style>
    [data-testid="stViewerBadge"] {
        display: none !important;
    }
    </style>
    """,
    unsafe_allow_html=True,
)

# =========================================================
# STYLING
# =========================================================

st.markdown(
    """
    <style>
    .stApp {
        background-color: #F8F9FC;
    }

    .block-container {
        padding-top: 1.7rem;
        padding-bottom: 2.5rem;
        max-width: 1200px;
    }

    h1, h2, h3 {
        color: #283653;
    }

    div[data-testid="stAlert"] {
        border-radius: 12px;
    }

    div[data-testid="stFileUploader"] {
        border-radius: 12px;
    }

    div.stButton > button,
    div.stDownloadButton > button {
        border-radius: 9px;
        font-weight: 600;
    }

    .app-subtitle {
        color: #68758B;
        font-size: 1.05rem;
        margin-top: -8px;
    }
    </style>
    """,
    unsafe_allow_html=True,
)

# =========================================================
# SESSION STATE
# =========================================================

DEFAULTS = {
    "notes_input": "",
    "pending_notes_input": None,
    "meeting_result": None,
    "image_extracted_text": "",
    "audio_extracted_text": "",
}

for key, value in DEFAULTS.items():
    if key not in st.session_state:
        st.session_state[key] = value

# Apply extracted text BEFORE creating the notes widget.
# This avoids Streamlit's session-state modification error.
if st.session_state.pending_notes_input is not None:
    st.session_state.notes_input = st.session_state.pending_notes_input
    st.session_state.pending_notes_input = None

# =========================================================
# GROQ CLIENT
# =========================================================

def get_client():
    api_key = os.getenv("GROQ_API_KEY")

    if not api_key:
        raise ValueError(
            "GROQ_API_KEY was not found. Check your existing .env file."
        )

    return Groq(api_key=api_key)


# =========================================================
# HANDWRITTEN / PRINTED IMAGE EXTRACTION
# =========================================================

def read_image_with_ai(uploaded_file):
    import base64

    client = get_client()
    image_bytes = uploaded_file.getvalue()

    extension = uploaded_file.name.lower().rsplit(".", 1)[-1]

    mime_types = {
        "png": "image/png",
        "jpg": "image/jpeg",
        "jpeg": "image/jpeg",
        "webp": "image/webp",
    }

    mime_type = mime_types.get(extension, "image/png")
    encoded_image = base64.b64encode(image_bytes).decode("utf-8")

    response = client.chat.completions.create(
        model=VISION_MODEL,
        messages=[
            {
                "role": "system",
                "content": (
                    "You are an accurate meeting-notes transcription assistant. "
                    "Read handwritten or printed notes from the image. "
                    "Preserve names, dates, deadlines, numbers, decisions, "
                    "and task assignments. Do not invent missing information. "
                    "Mark unreadable text as [unclear]. Return only the "
                    "transcribed notes with line breaks."
                ),
            },
            {
                "role": "user",
                "content": [
                    {
                        "type": "text",
                        "text": (
                            "Transcribe all meeting notes visible in this image. "
                            "Preserve the original meaning and distinguish "
                            "headings, tasks, names, and dates where possible."
                        ),
                    },
                    {
                        "type": "image_url",
                        "image_url": {
                            "url": (
                                f"data:{mime_type};base64,{encoded_image}"
                            )
                        },
                    },
                ],
            },
        ],
        temperature=0.1,
        max_tokens=2500,
    )

    return response.choices[0].message.content.strip()


# =========================================================
# MULTILINGUAL AUDIO TRANSCRIPTION
# =========================================================

def transcribe_audio_bytes(audio_bytes, filename):
    client = get_client()

    transcription_prompt = (
        "Transcribe the speech as spoken, preserving the original language "
        "and wording. The recording may contain English, Hindi, Hinglish, "
        "or frequent switching between Hindi and English within the same "
        "sentence. Do not translate Hindi into English or English into Hindi. "
        "Use Devanagari for Hindi words when appropriate and Latin script "
        "for English words and technical terms when appropriate. Preserve "
        "names, numbers, dates, deadlines, and business terminology. "
        "Common workplace terms may include meeting, deadline, website, "
        "design, payment gateway, testing, budget, approval, client, "
        "marketing, deployment, and follow-up. Do not invent speech that "
        "is not audible."
    )

    response = client.audio.transcriptions.create(
        file=(filename, audio_bytes),
        model=AUDIO_MODEL,
        prompt=transcription_prompt,
        response_format="text",
    )

    if isinstance(response, str):
        return response.strip()

    return str(response).strip()


# =========================================================
# VIDEO AUDIO EXTRACTION
# =========================================================

def transcribe_video(uploaded_file):
    video_path = None
    audio_path = None

    try:
        suffix = os.path.splitext(uploaded_file.name)[1] or ".mp4"

        with tempfile.NamedTemporaryFile(
            delete=False,
            suffix=suffix,
        ) as video_temp:
            video_temp.write(uploaded_file.getvalue())
            video_path = video_temp.name

        with tempfile.NamedTemporaryFile(
            delete=False,
            suffix=".mp3",
        ) as audio_temp:
            audio_path = audio_temp.name

        try:
            from moviepy import VideoFileClip
        except ImportError:
            from moviepy.editor import VideoFileClip

        with VideoFileClip(video_path) as clip:
            if clip.audio is None:
                raise ValueError(
                    "No audio track was found in this video."
                )

            clip.audio.write_audiofile(
                audio_path,
                logger=None,
            )

        with open(audio_path, "rb") as audio_file:
            audio_bytes = audio_file.read()

        return transcribe_audio_bytes(
            audio_bytes,
            "meeting_audio.mp3",
        )

    finally:
        for path in (video_path, audio_path):
            if path and os.path.exists(path):
                try:
                    os.remove(path)
                except OSError:
                    pass


# =========================================================
# HINGLISH-AWARE MEETING ANALYSIS
# =========================================================

def analyze_meeting(notes):
    if not notes or not notes.strip():
        raise ValueError(
            "Please enter meeting notes before analyzing."
        )

    client = get_client()

    system_prompt = """
You are MeetBrief AI, a professional meeting-notes analyst.

The input may be typed notes, a transcript, or speech transcribed from
an audio or video recording.

The content may be in English, Hindi, Hinglish, or a mixture of languages.

Understand Hindi and Hinglish instructions and business terminology
without requiring the input to be translated first.

Return ONLY valid JSON. Do not use Markdown code fences.

Use exactly this structure:

{
  "meeting_title": "A concise meeting title in English",
  "summary": "A concise professional summary in English",
  "key_decisions": [
    "Decision 1"
  ],
  "action_items": [
    {
      "task": "Specific task in English",
      "owner": "Person's name or Unassigned",
      "due_date": "Date stated in notes or Not specified",
      "priority": "High, Medium, Low, or Not specified",
      "status": "Pending",
      "source": "Brief supporting phrase from the notes"
    }
  ],
  "open_questions": [
    "Question requiring clarification"
  ],
  "risks_blockers": [
    "Risk, dependency, or blocker"
  ]
}

RULES:

1. Generate the meeting title, summary, decisions, action items,
   open questions, and risks in clear professional English.

2. Understand Hindi and Hinglish instructions accurately.
   For example, "Friday tak complete kar dena" means the task
   should be completed by Friday.

3. Preserve names, company names, technical terms, numbers,
   dates, amounts, and deadlines accurately.

4. Never invent owners, deadlines, decisions, or commitments.

5. If an owner is missing, use "Unassigned".

6. If a deadline is missing, use "Not specified".

7. Include action items only when a task or follow-up is indicated.

8. Record unresolved approvals and dependencies as open questions
   or blockers when appropriate.

9. If a category has no items, return an empty list.

10. If the transcript contains an unclear or potentially incorrect
    word, do not guess its meaning when that could change an action item.

11. The original transcript and user-provided notes must remain unchanged.
    Do not translate or overwrite the original input while creating the brief.

12. If an instruction is ambiguous because of transcription errors
    or mixed-language wording, flag the uncertainty rather than guessing.

13. Return valid JSON only.
"""

    response = client.chat.completions.create(
        model=TEXT_MODEL,
        messages=[
            {
                "role": "system",
                "content": system_prompt,
            },
            {
                "role": "user",
                "content": (
                    "Create a meeting brief from these notes. "
                    "The notes may contain Hindi, English, or Hinglish. "
                    "Understand their meaning and produce the structured "
                    "meeting brief in English.\n\n"
                    + notes
                ),
            },
        ],
        temperature=0.2,
        max_tokens=3500,
    )

    content = response.choices[0].message.content.strip()

    if content.startswith("```"):
        content = content.replace("```json", "", 1)
        content = content.replace("```", "").strip()

    start = content.find("{")
    end = content.rfind("}")

    if start == -1 or end == -1:
        raise ValueError(
            "The AI did not return valid JSON. Please try again."
        )

    result = json.loads(content[start:end + 1])

    result.setdefault("meeting_title", "Meeting Brief")
    result.setdefault("summary", "")
    result.setdefault("key_decisions", [])
    result.setdefault("action_items", [])
    result.setdefault("open_questions", [])
    result.setdefault("risks_blockers", [])

    return result


# =========================================================
# EXPORT HELPERS
# =========================================================

def build_action_dataframe(result):
    columns = [
        "Task",
        "Owner",
        "Due Date",
        "Priority",
        "Status",
        "Source",
    ]

    rows = []

    for item in result.get("action_items", []):
        rows.append(
            {
                "Task": item.get("task", ""),
                "Owner": item.get("owner", "Unassigned"),
                "Due Date": item.get("due_date", "Not specified"),
                "Priority": item.get("priority", "Not specified"),
                "Status": item.get("status", "Pending"),
                "Source": item.get("source", ""),
            }
        )

    return pd.DataFrame(rows, columns=columns)


def build_docx(result):
    document = Document()

    document.add_heading(
        result.get("meeting_title", "Meeting Brief"),
        level=0,
    )

    document.add_paragraph(
        "Generated by MeetBrief AI | "
        + datetime.now().strftime("%d %B %Y, %I:%M %p")
    )

    document.add_heading("Executive Summary", level=1)
    document.add_paragraph(result.get("summary", ""))

    document.add_heading("Key Decisions", level=1)
    decisions = result.get("key_decisions", [])

    if decisions:
        for decision in decisions:
            document.add_paragraph(str(decision), style="List Bullet")
    else:
        document.add_paragraph("No key decisions recorded.")

    document.add_heading("Action Items", level=1)
    actions = result.get("action_items", [])

    if actions:
        table = document.add_table(rows=1, cols=5)
        table.style = "Light Shading Accent 1"

        headers = ["Task", "Owner", "Due Date", "Priority", "Status"]

        for cell, header in zip(table.rows[0].cells, headers):
            cell.text = header

        for item in actions:
            cells = table.add_row().cells

            values = [
                item.get("task", ""),
                item.get("owner", "Unassigned"),
                item.get("due_date", "Not specified"),
                item.get("priority", "Not specified"),
                item.get("status", "Pending"),
            ]

            for cell, value in zip(cells, values):
                cell.text = str(value)

    else:
        document.add_paragraph("No action items identified.")

    for title, key in [
        ("Open Questions", "open_questions"),
        ("Risks and Blockers", "risks_blockers"),
    ]:
        document.add_heading(title, level=1)
        entries = result.get(key, [])

        if entries:
            for entry in entries:
                document.add_paragraph(str(entry), style="List Bullet")
        else:
            document.add_paragraph("None identified.")

    buffer = io.BytesIO()
    document.save(buffer)
    buffer.seek(0)

    return buffer.getvalue()


def build_pdf(result):
    buffer = io.BytesIO()

    pdf = SimpleDocTemplate(
        buffer,
        pagesize=A4,
        rightMargin=0.55 * inch,
        leftMargin=0.55 * inch,
        topMargin=0.6 * inch,
        bottomMargin=0.6 * inch,
    )

    styles = getSampleStyleSheet()

    title_style = ParagraphStyle(
        "MeetBriefTitle",
        parent=styles["Title"],
        textColor=colors.HexColor("#283653"),
        spaceAfter=12,
    )

    heading_style = ParagraphStyle(
        "MeetBriefHeading",
        parent=styles["Heading2"],
        textColor=colors.HexColor("#465B85"),
        spaceBefore=14,
        spaceAfter=6,
    )

    body_style = styles["BodyText"]

    def safe_text(value):
        text_value = str(value or "")
        return (
            text_value.replace("&", "&amp;")
            .replace("<", "&lt;")
            .replace(">", "&gt;")
            .replace("\n", "<br/>")
        )

    story = [
        Paragraph(
            safe_text(result.get("meeting_title", "Meeting Brief")),
            title_style,
        ),
        Paragraph(
            "Generated by MeetBrief AI | "
            + datetime.now().strftime("%d %B %Y, %I:%M %p"),
            body_style,
        ),
        Spacer(1, 12),
        Paragraph("Executive Summary", heading_style),
        Paragraph(safe_text(result.get("summary", "")), body_style),
    ]

    for title, key in [
        ("Key Decisions", "key_decisions"),
        ("Open Questions", "open_questions"),
        ("Risks and Blockers", "risks_blockers"),
    ]:
        story.append(Paragraph(title, heading_style))
        entries = result.get(key, [])

        if entries:
            for entry in entries:
                story.append(
                    Paragraph(
                        "&#8226; " + safe_text(entry),
                        body_style,
                    )
                )
                story.append(Spacer(1, 4))
        else:
            story.append(Paragraph("None identified.", body_style))

    story.append(Paragraph("Action Items", heading_style))
    actions = result.get("action_items", [])

    if actions:
        data = [[
            "Task",
            "Owner",
            "Due Date",
            "Priority",
            "Status",
        ]]

        for item in actions:
            data.append([
                Paragraph(safe_text(item.get("task", "")), body_style),
                Paragraph(safe_text(item.get("owner", "Unassigned")), body_style),
                Paragraph(safe_text(item.get("due_date", "Not specified")), body_style),
                Paragraph(safe_text(item.get("priority", "Not specified")), body_style),
                Paragraph(safe_text(item.get("status", "Pending")), body_style),
            ])

        table = Table(
            data,
            colWidths=[
                2.05 * inch,
                1.05 * inch,
                1.05 * inch,
                0.8 * inch,
                0.85 * inch,
            ],
            repeatRows=1,
        )

        table.setStyle(
            TableStyle([
                (
                    "BACKGROUND",
                    (0, 0),
                    (-1, 0),
                    colors.HexColor("#E7ECF7"),
                ),
                (
                    "TEXTCOLOR",
                    (0, 0),
                    (-1, 0),
                    colors.HexColor("#283653"),
                ),
                (
                    "GRID",
                    (0, 0),
                    (-1, -1),
                    0.4,
                    colors.HexColor("#D5DCE8"),
                ),
                ("VALIGN", (0, 0), (-1, -1), "TOP"),
                ("LEFTPADDING", (0, 0), (-1, -1), 5),
                ("RIGHTPADDING", (0, 0), (-1, -1), 5),
                ("TOPPADDING", (0, 0), (-1, -1), 6),
                ("BOTTOMPADDING", (0, 0), (-1, -1), 6),
            ])
        )

        story.append(table)

    else:
        story.append(
            Paragraph("No action items identified.", body_style)
        )

    pdf.build(story)
    buffer.seek(0)

    return buffer.getvalue()


# =========================================================
# HEADER AND SIDEBAR
# =========================================================

st.title("📝 MeetBrief AI")

st.markdown(
    '<p class="app-subtitle">'
    "Turn meeting notes and recordings into clear summaries, "
    "decisions, and actionable tasks."
    "</p>",
    unsafe_allow_html=True,
)

with st.sidebar:
    st.header("About MeetBrief AI")

    st.write(
        "MeetBrief AI converts meeting information into structured, "
        "easy-to-review meeting briefs."
    )

    st.divider()

    st.subheader("Supported inputs")
    st.write("• Typed notes and transcripts")
    st.write("• Handwritten or printed note images")
    st.write("• Audio recordings")
    st.write("• Video recordings with audio")
    st.write("• English, Hindi, and Hinglish speech")

    st.divider()

    st.caption(
        "AI-generated content can contain errors. Review names, "
        "dates, deadlines, and task assignments before sharing."
    )


# =========================================================
# INPUT TABS
# =========================================================

text_tab, image_tab, recording_tab = st.tabs(
    [
        "✍️ Text / Transcript",
        "📷 Handwritten Notes Image",
        "🎙️ Audio / Video Recording",
    ]
)

with text_tab:
    st.subheader("Meeting notes")

    st.caption(
        "Paste meeting notes or a transcript in English, Hindi, "
        "or Hinglish. Extracted text can also be reviewed here."
    )

    st.text_area(
        "Meeting notes or transcript",
        height=260,
        placeholder=(
            "Example:\n"
            "Meeting: Website launch planning\n"
            "Priya homepage ka design October 14 tak complete karegi.\n"
            "Rahul will start payment gateway testing by October 16."
        ),
        key="notes_input",
        label_visibility="collapsed",
    )

    if st.button("Clear meeting notes", key="clear_notes"):
        st.session_state.pending_notes_input = ""
        st.session_state.meeting_result = None
        st.rerun()


with image_tab:
    st.info(
        "Upload a photo or scan of handwritten or printed meeting notes. "
        "Handwriting recognition may be imperfect; review the extracted "
        "text before generating the brief."
    )

    image_file = st.file_uploader(
        "Choose an image from your device",
        type=["png", "jpg", "jpeg", "webp"],
        key="notes_image",
    )

    if image_file and st.button(
        "Extract text from image",
        key="extract_image",
        type="primary",
    ):
        with st.spinner("Reading the image with AI..."):
            try:
                extracted = read_image_with_ai(image_file)

                st.session_state.image_extracted_text = extracted
                st.session_state.pending_notes_input = extracted

                st.success(
                    "Text extracted! Switch to Text / Transcript "
                    "to review and edit the extracted notes."
                )

                st.text_area(
                    "Extracted text preview",
                    value=extracted,
                    height=180,
                    key="image_preview",
                )

            except Exception as exc:
                st.error(f"Could not read the image: {exc}")

    elif st.session_state.image_extracted_text:
        st.text_area(
            "Previously extracted text",
            value=st.session_state.image_extracted_text,
            height=180,
            key="image_preview_saved",
        )


with recording_tab:
    st.info(
        "Upload an audio recording or a video containing speech. "
        "Transcription supports multiple languages, including Hindi "
        "and English. Hindi-English code-switching (Hinglish) may work "
        "well, but recognition can vary. The transcript aims to preserve "
        "the original spoken language rather than translate it. "
        "Review the transcript before analysis."
    )

    audio_file = st.file_uploader(
        "Upload an audio recording",
        type=[
            "mp3", "mp4", "mpeg", "mpga",
            "m4a", "wav", "webm", "ogg", "flac",
        ],
        key="meeting_audio",
    )

    if audio_file and st.button(
        "Transcribe audio",
        key="transcribe_audio",
        type="primary",
    ):
        with st.spinner("Transcribing audio..."):
            try:
                transcript = transcribe_audio_bytes(
                    audio_file.getvalue(),
                    audio_file.name,
                )

                st.session_state.audio_extracted_text = transcript
                st.session_state.pending_notes_input = transcript

                st.success(
                    "Audio transcribed! Switch to Text / Transcript "
                    "to review the original-language transcript."
                )

                st.text_area(
                    "Audio transcript preview",
                    value=transcript,
                    height=180,
                    key="audio_preview",
                )

            except Exception as exc:
                st.error(f"Could not transcribe audio: {exc}")

    elif st.session_state.audio_extracted_text:
        st.text_area(
            "Previously transcribed audio",
            value=st.session_state.audio_extracted_text,
            height=180,
            key="audio_preview_saved",
        )

    st.divider()

    video_file = st.file_uploader(
        "Upload a video recording",
        type=["mp4", "mov", "mkv", "avi", "webm"],
        key="meeting_video",
    )

    if video_file and st.button(
        "Extract and transcribe video audio",
        key="transcribe_video",
        type="primary",
    ):
        with st.spinner(
            "Extracting audio from the video and transcribing it..."
        ):
            try:
                transcript = transcribe_video(video_file)

                st.session_state.audio_extracted_text = transcript
                st.session_state.pending_notes_input = transcript

                st.success(
                    "Video transcribed! Switch to Text / Transcript "
                    "to review the original-language transcript."
                )

                st.text_area(
                    "Video transcript preview",
                    value=transcript,
                    height=180,
                    key="video_preview",
                )

            except Exception as exc:
                st.error(
                    "Could not process the video. Check that it contains "
                    f"audio and that video dependencies are installed. "
                    f"Details: {exc}"
                )


# =========================================================
# START ANALYSIS
# =========================================================

st.divider()
st.subheader("Generate your meeting brief")

if st.button(
    "✨ Start analysis",
    type="primary",
    use_container_width=True,
):
    notes = st.session_state.notes_input.strip()

    if not notes:
        st.warning(
            "Please enter notes or extract text from an image "
            "or recording before starting the analysis."
        )

    else:
        with st.spinner(
            "Analyzing the meeting and organizing the action items..."
        ):
            try:
                result = analyze_meeting(notes)
                st.session_state.meeting_result = result
                st.success("Your meeting brief is ready!")

            except Exception as exc:
                st.error(f"Analysis failed: {exc}")


# =========================================================
# DISPLAY RESULTS
# =========================================================

result = st.session_state.meeting_result

if result:
    st.divider()
    st.header(result.get("meeting_title", "Meeting Brief"))

    action_items = result.get("action_items", [])
    decisions = result.get("key_decisions", [])

    metric1, metric2, metric3 = st.columns(3)

    metric1.metric("Action Items", len(action_items))
    metric2.metric("Key Decisions", len(decisions))
    metric3.metric(
        "Open Questions",
        len(result.get("open_questions", [])),
    )

    st.subheader("Executive Summary")
    st.write(result.get("summary", "No summary available."))

    st.subheader("Key Decisions")

    if decisions:
        for decision in decisions:
            st.markdown(f"- {decision}")
    else:
        st.write("No key decisions identified.")

    st.subheader("Action Items")

    if action_items:
        action_df = build_action_dataframe(result)

        st.dataframe(
            action_df,
            use_container_width=True,
            hide_index=True,
        )
    else:
        st.info("No action items were identified in the notes.")

    with st.expander("Open Questions", expanded=True):
        questions = result.get("open_questions", [])

        if questions:
            for question in questions:
                st.markdown(f"- {question}")
        else:
            st.write("No open questions identified.")

    with st.expander("Risks and Blockers", expanded=True):
        risks = result.get("risks_blockers", [])

        if risks:
            for risk in risks:
                st.markdown(f"- {risk}")
        else:
            st.write("No risks or blockers identified.")

    # =====================================================
    # EXPORTS
    # =====================================================

    st.divider()
    st.subheader("Export your meeting brief")

    export1, export2, export3 = st.columns(3)

    safe_title = "".join(
        character
        for character in result.get("meeting_title", "meeting_brief")
        if character.isalnum() or character in (" ", "-", "_")
    ).strip().replace(" ", "_")

    safe_title = safe_title or "meeting_brief"

    with export1:
        try:
            docx_bytes = build_docx(result)

            st.download_button(
                "📄 Download Word",
                data=docx_bytes,
                file_name=f"{safe_title}.docx",
                mime=(
                    "application/vnd.openxmlformats-officedocument."
                    "wordprocessingml.document"
                ),
                use_container_width=True,
            )
        except Exception as exc:
            st.error(f"Word export failed: {exc}")

    with export2:
        try:
            pdf_bytes = build_pdf(result)

            st.download_button(
                "📕 Download PDF",
                data=pdf_bytes,
                file_name=f"{safe_title}.pdf",
                mime="application/pdf",
                use_container_width=True,
            )
        except Exception as exc:
            st.error(f"PDF export failed: {exc}")

    with export3:
        action_csv = build_action_dataframe(result).to_csv(index=False)

        st.download_button(
            "📊 Download Tasks CSV",
            data=action_csv,
            file_name=f"{safe_title}_action_items.csv",
            mime="text/csv",
            use_container_width=True,
        )
